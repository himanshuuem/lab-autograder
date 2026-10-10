"""Build .docx files the way students write them in Word."""
import io

import docx
from docx.shared import Inches, Pt


def make_docx(path, blocks) -> str:
    """blocks: ("h", text) | ("p", text) | ("code", text, font) | ("img", png_bytes, width_in)"""
    d = docx.Document()
    for b in blocks:
        if b[0] == "h":
            d.add_heading(b[1], level=2)
        elif b[0] == "p":
            d.add_paragraph(b[1])
        elif b[0] == "code":
            for line in b[1].split("\n"):
                para = d.add_paragraph()
                para.paragraph_format.space_after = Pt(0)
                run = para.add_run(line)
                run.font.name = b[2]
                run.font.size = Pt(10)
        elif b[0] == "img":
            d.add_picture(io.BytesIO(b[1]), width=Inches(b[2] if len(b) > 2 else 4.5))
    d.save(path)
    return str(path)
